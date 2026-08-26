#!/usr/bin/env python3
# =============================================================================
# CHECK DUPLICATES
# =============================================================================
# Generated: 2026-08-26 10:20 AM EDT
#
# WHAT THIS SCRIPT DOES
# ----------------------------------------------------------------------------
# Runs BEFORE build_issue_csv.py, against the same DataWaiting folder.
# Looks inside every tifflist and Town List file there for duplicate
# entries WITHIN that one file - the same TIFF filename listed twice in
# a tifflist, or the same paper number appearing on more than one row
# (or, for the "transposed" Town List layout, more than one column) of
# a Town List. This is different from - and doesn't replace - the
# cross-file duplicate detection build_issue_csv.py already does when
# the SAME paper number shows up in two different files; that part was
# already handled and is untouched.
#
# If a file has no internal duplicates, it's left completely alone.
#
# If a file DOES have internal duplicates:
#   1. The original file is renamed with "-has_duplicates" added before
#      its extension, and moved straight to DataFilesCompleted - so
#      anyone browsing that folder sees immediately that something
#      needed fixing, and can open the file to see exactly what the
#      raw data looked like before the fix.
#   2. A corrected copy - first occurrence of each duplicate kept, the
#      repeats removed - is written back into DataWaiting under the
#      file's ORIGINAL name, so build_issue_csv.py picks it up on its
#      normal run without needing to know anything happened at all.
#
# This script is intentionally independent of build_issue_csv.py - it
# has its own (simpler) copy of the file-recognition and Town-List-
# layout logic below, rather than importing anything from that file.
# That's a deliberate choice, not an oversight: it keeps this
# duplicate-fixing step free to change on its own without any risk of
# affecting the main script's behavior, and vice versa. It only needs
# to find each row/column's paper number well enough to spot a repeat -
# it never needs to fully parse town names the way build_issue_csv.py
# does, since that's not this script's job.
#
# HOW TO RUN
# ----------------------------------------------------------------------------
#   python check_duplicates.py <DataWaiting_dir> <DataFilesCompleted_dir>
#
# Logs a line to <root>/_Logs/run_log.txt (the SAME run log
# build_issue_csv.py uses - "root" is DataWaiting's parent folder) for
# every file it fixes, plus one summary line every run, fixed or not -
# so there's always a record this step ran at all.
# =============================================================================

import os
import re
import sys
import csv
from datetime import datetime

DUPLICATE_TAG = "-has_duplicates"

# ---- File recognition (independent copy - see module note above) ----------

TIFF_NAME_SEARCH_PATTERN = re.compile(
    r'(\d{8}_v[A-Za-z0-9]+_n\d+(?:[_-]\d+)?\.tiff?)',
    re.IGNORECASE,
)
TIFFLIST_FILENAME_PATTERN = re.compile(
    r'^(.*?)_?(?i:tifflist)\.(?i:txt|csv)$'
)
TOWNLIST_KEYWORD_PATTERN = re.compile(r'(?i:town[_ -]?list)')
ISSUE_PREFIX_PATTERN = re.compile(r'^(\d{8})_v([A-Za-z0-9]+)_n(\d+)')


def read_text_lines(path):
    """Same UTF-8 / UTF-16 tolerance as build_issue_csv.py's reader, so
    a file written by a Windows 'dir' listing (PowerShell's default
    UTF-16 output) doesn't look unreadable here either."""
    try:
        with open(path, "r", encoding="utf-8-sig") as f:
            return f.readlines()
    except UnicodeDecodeError:
        with open(path, "r", encoding="utf-16") as f:
            return f.readlines()


def is_town_list_filename(filename):
    if not filename.lower().endswith((".xlsx", ".csv")):
        return False
    return bool(TOWNLIST_KEYWORD_PATTERN.search(filename))


def file_contains_tiff_filenames(path):
    """Content-sniffing fallback for a tifflist file whose name doesn't
    follow the usual convention - same 50%-of-lines-match rule
    build_issue_csv.py uses."""
    try:
        lines = [line for line in read_text_lines(path) if line.strip()]
    except (UnicodeDecodeError, OSError):
        return False
    if not lines:
        return False
    matching = sum(1 for line in lines if TIFF_NAME_SEARCH_PATTERN.search(line))
    return matching / len(lines) >= 0.5


def classify_input_files(directory):
    """
    Returns (tiff_filenames, townlist_filenames) - the same recognition
    rules build_issue_csv.py's find_tiff_source_files() /
    find_town_list_source_files() use, reimplemented independently here
    (see the module note at the top of this file). Files already tagged
    "-has_duplicates" are skipped defensively - they should never
    actually be sitting in DataWaiting (they get moved out immediately),
    but there's no reason to risk re-processing one if it somehow is.
    """
    tiff_by_name = []
    maybe_by_content = []
    townlist_filenames = []

    for f in sorted(os.listdir(directory)):
        if f.startswith("."):
            continue
        if DUPLICATE_TAG in f:
            continue
        path = os.path.join(directory, f)
        if not os.path.isfile(path):
            continue

        if is_town_list_filename(f):
            townlist_filenames.append(f)
        elif TIFFLIST_FILENAME_PATTERN.match(f):
            tiff_by_name.append(f)
        elif f.lower().endswith((".txt", ".csv")):
            maybe_by_content.append(f)

    tiff_by_content = [
        f for f in maybe_by_content
        if file_contains_tiff_filenames(os.path.join(directory, f))
    ]

    return sorted(tiff_by_name + tiff_by_content), sorted(townlist_filenames)


# ---- Tifflist duplicate detection ------------------------------------------

def find_tifflist_duplicates(path):
    """
    Returns (fixed_lines, duplicate_count). fixed_lines is the file's
    lines with every REPEAT of an already-seen TIFF filename removed
    (first occurrence kept, same "keep the first one seen" rule used
    everywhere else in this pipeline); duplicate_count is how many
    lines were removed. If duplicate_count is 0, fixed_lines is just
    the file's original lines, unchanged.
    """
    lines = read_text_lines(path)
    seen = set()
    fixed_lines = []
    duplicate_count = 0

    for raw_line in lines:
        match = TIFF_NAME_SEARCH_PATTERN.search(raw_line)
        if match:
            filename = match.group(1)
            if filename in seen:
                duplicate_count += 1
                continue
            seen.add(filename)
        fixed_lines.append(raw_line)

    return fixed_lines, duplicate_count


# ---- Town List duplicate detection -----------------------------------------
# A lightweight, standalone version of build_issue_csv.py's three-layout
# detection - just enough to find each row's (or, for the transposed
# layout, each column's) paper number, not to fully parse town names.

def read_grid(path):
    extension = os.path.splitext(path)[1].lower()
    if extension == ".xlsx":
        import openpyxl
        workbook = openpyxl.load_workbook(path, data_only=True)
        worksheet = workbook.active
        grid = []
        for row in worksheet.iter_rows(values_only=True):
            text_row = []
            for cell in row:
                if cell is None:
                    text_row.append("")
                elif isinstance(cell, float) and cell.is_integer():
                    text_row.append(str(int(cell)))
                else:
                    text_row.append(str(cell).strip())
            grid.append(text_row)
        return grid
    elif extension == ".csv":
        with open(path, "r", newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            return [[cell.strip() for cell in row] for row in reader]
    else:
        raise ValueError(f"Unexpected Town List extension: {extension}")


def write_grid(path, grid):
    extension = os.path.splitext(path)[1].lower()
    if extension == ".xlsx":
        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in grid:
            sheet.append(row)
        workbook.save(path)
    elif extension == ".csv":
        with open(path, "w", newline="", encoding="utf-8") as f:
            csv.writer(f, lineterminator="\n").writerows(grid)
    else:
        raise ValueError(f"Unexpected Town List extension: {extension}")


def fix_town_list_duplicates(grid):
    """
    Returns (fixed_grid, duplicate_count). Removes rows - or, for the
    transposed layout, columns - whose paper number repeats one already
    kept (first occurrence wins, same rule as everywhere else). If
    duplicate_count is 0, fixed_grid is the same grid, unchanged.
    """
    if not grid or not grid[0]:
        return grid, 0

    header_row = grid[0]
    header_lower = [c.strip().lower() for c in header_row]
    first_cell = header_row[0].lower() if header_row else ""
    is_slim = "subject" in header_lower and "geographic" in header_lower

    if first_cell.startswith("paper") and not is_slim:
        # Transposed format - paper numbers run across the header row,
        # so a duplicate means a repeated COLUMN.
        seen = set()
        keep_columns = [0]  # always keep the row-label column
        duplicate_count = 0
        for col_index in range(1, len(header_row)):
            text = header_row[col_index].strip()
            if text.isdigit() and text in seen:
                duplicate_count += 1
                continue
            if text.isdigit():
                seen.add(text)
            keep_columns.append(col_index)

        if duplicate_count == 0:
            return grid, 0

        fixed_grid = [
            [row[c] if c < len(row) else "" for c in keep_columns]
            for row in grid
        ]
        return fixed_grid, duplicate_count

    # Row-based formats (plain "one row per issue", or Slim) - a
    # duplicate means a repeated ROW.
    if is_slim:
        def row_key(row):
            if not row or not row[0].strip():
                return None
            match = ISSUE_PREFIX_PATTERN.match(row[0].strip())
            return (match.group(2), match.group(3)) if match else None
    else:
        paper_no_col = None
        for col_index, header_text in enumerate(header_row):
            cleaned = header_text.strip().lower().rstrip(".")
            if cleaned in ("no", "paper no", "papernum", "paper number"):
                paper_no_col = col_index
                break
        if paper_no_col is None:
            return grid, 0  # can't find a paper-number column - leave alone

        def row_key(row):
            if paper_no_col >= len(row):
                return None
            text = row[paper_no_col].strip()
            return text if text.isdigit() else None

    seen = set()
    fixed_rows = [header_row]
    duplicate_count = 0
    for row in grid[1:]:
        key = row_key(row)
        if key is not None:
            if key in seen:
                duplicate_count += 1
                continue
            seen.add(key)
        fixed_rows.append(row)

    if duplicate_count == 0:
        return grid, 0
    return fixed_rows, duplicate_count


# ---- Shared run-log helper (independent copy - see module note) -----------

def append_run_log(waiting_dir, message):
    """Appends one line to <root>/_Logs/run_log.txt, where root is
    DataWaiting's parent folder - the SAME log file build_issue_csv.py
    writes to, so both scripts' activity shows up in one place."""
    log_root = os.path.dirname(os.path.abspath(waiting_dir.rstrip(os.sep)))
    logs_dir = os.path.join(log_root, "_Logs")
    os.makedirs(logs_dir, exist_ok=True)
    log_path = os.path.join(logs_dir, "run_log.txt")
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(f"{timestamp}  {message}\n")


# ---- Fixing a single file --------------------------------------------------

def make_tagged_name(filename):
    base, ext = os.path.splitext(filename)
    return f"{base}{DUPLICATE_TAG}{ext}"


def process_tifflist_file(filename, waiting_dir, completed_dir, log_lines):
    path = os.path.join(waiting_dir, filename)
    try:
        fixed_lines, duplicate_count = find_tifflist_duplicates(path)
    except (UnicodeDecodeError, OSError) as e:
        log_lines.append(f"[{filename}] couldn't be read ({e}) - left as-is.")
        return

    if duplicate_count == 0:
        return

    tagged_name = make_tagged_name(filename)
    os.rename(path, os.path.join(completed_dir, tagged_name))

    with open(os.path.join(waiting_dir, filename), "w", encoding="utf-8") as f:
        f.writelines(fixed_lines)

    log_lines.append(
        f"[duplicate_fix] {filename} -> {DATA_COMPLETED_LABEL}/{tagged_name} "
        f"({duplicate_count} duplicate line(s) removed; corrected copy "
        f"kept in DataWaiting as {filename})"
    )
    print(f"  fixed duplicates in '{filename}' ({duplicate_count} removed) - "
          f"original moved to DataFilesCompleted as '{tagged_name}'")


def process_town_list_file(filename, waiting_dir, completed_dir, log_lines):
    path = os.path.join(waiting_dir, filename)
    try:
        grid = read_grid(path)
        fixed_grid, duplicate_count = fix_town_list_duplicates(grid)
    except Exception as e:
        log_lines.append(f"[{filename}] couldn't be read ({e}) - left as-is.")
        return

    if duplicate_count == 0:
        return

    tagged_name = make_tagged_name(filename)
    os.rename(path, os.path.join(completed_dir, tagged_name))

    write_grid(os.path.join(waiting_dir, filename), fixed_grid)

    log_lines.append(
        f"[duplicate_fix] {filename} -> {DATA_COMPLETED_LABEL}/{tagged_name} "
        f"({duplicate_count} duplicate row/column(s) removed; corrected "
        f"copy kept in DataWaiting as {filename})"
    )
    print(f"  fixed duplicates in '{filename}' ({duplicate_count} removed) - "
          f"original moved to DataFilesCompleted as '{tagged_name}'")


DATA_COMPLETED_LABEL = "DataFilesCompleted"


def main():
    if len(sys.argv) != 3:
        print("Usage: python check_duplicates.py <DataWaiting_dir> <DataFilesCompleted_dir>")
        sys.exit(1)

    waiting_dir, completed_dir = sys.argv[1], sys.argv[2]
    if not os.path.isdir(waiting_dir):
        print(f"ERROR: '{waiting_dir}' is not a directory.")
        sys.exit(1)
    os.makedirs(completed_dir, exist_ok=True)

    tiff_filenames, townlist_filenames = classify_input_files(waiting_dir)
    print(f"Checking {len(tiff_filenames)} tifflist file(s) and "
          f"{len(townlist_filenames)} Town List file(s) for internal duplicates...")

    log_lines = []
    for filename in tiff_filenames:
        process_tifflist_file(filename, waiting_dir, completed_dir, log_lines)
    for filename in townlist_filenames:
        process_town_list_file(filename, waiting_dir, completed_dir, log_lines)

    for line in log_lines:
        append_run_log(waiting_dir, line)

    if log_lines:
        print()
        print(f"Fixed {len(log_lines)} file(s) with internal duplicates.")
        append_run_log(waiting_dir, f"check_duplicates: fixed {len(log_lines)} file(s)")
    else:
        print("No internal duplicates found.")
        append_run_log(waiting_dir, "check_duplicates: no internal duplicates found")


if __name__ == "__main__":
    main()
